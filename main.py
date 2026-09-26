"""
Zm@il - application locale d'envoi de campagnes e-mail.

Backend FastAPI : rendu des modèles d'e-mail (Jinja2), envoi SMTP (TLS),
gestion des contacts et de la configuration, stockage en fichiers JSON.
"""

from __future__ import annotations

import asyncio
import base64
import csv
import io
import json
import mimetypes
import os
import re
import smtplib
import ssl
import sys
import time
import unicodedata
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid
from pathlib import Path
from typing import Any

from fastapi import Body, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from jinja2 import Environment, FileSystemLoader, TemplateNotFound, select_autoescape
from markupsafe import Markup, escape
from pydantic import BaseModel, Field

from constants import SOCIAL_ICON_SIZE, SOCIAL_ICON_URLS
from i18n import (
    DEFAULT_GROUP,
    DEFAULT_LANGUAGE,
    LANGUAGES,
    LANGUAGE_LABELS,
    TEMPLATE_ORDER,
    email_strings,
    normalize_language,
    sample_content,
    template_meta,
    tr,
)

# --------------------------------------------------------------------------- #
# Chemins & constantes
# --------------------------------------------------------------------------- #

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
UPLOADS_DIR = STATIC_DIR / "uploads"
TEMPLATES_DIR = BASE_DIR / "templates"
EMAIL_TEMPLATES_DIR = TEMPLATES_DIR / "email_templates"
DATA_DIR = BASE_DIR / "data"
CONTACTS_FILE = DATA_DIR / "contacts.json"
CONFIG_FILE = DATA_DIR / "config.json"

HOST = os.environ.get("EMAIL_SENDER_HOST", "127.0.0.1")
PORT = int(os.environ.get("EMAIL_SENDER_PORT", "8000"))

MASKED_PASSWORD = "••••••••"
BULK_DELAY_SECONDS = 1.0  # pause entre deux envois (limite les filtres anti-spam)
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
DEFAULT_ACCENT = "#2563eb"
EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[A-Za-z]{2,}$")

# En-tête envoyé par l'interface pour appliquer la langue avant même son
# enregistrement dans data/config.json (changement de langue instantané).
LANGUAGE_HEADER = "X-Language"


def sample_variables(lang: str = DEFAULT_LANGUAGE) -> dict[str, str]:
    """Contenu d'exemple des aperçus, dans la langue demandée."""
    return dict(sample_content(lang), accent_color=DEFAULT_ACCENT)

# --------------------------------------------------------------------------- #
# Dossiers & fichiers de données
# --------------------------------------------------------------------------- #


def default_config() -> dict[str, Any]:
    return {
        "smtp_server": "",
        "smtp_port": 587,
        "email": "",
        "password": "",
        "display_name": "",
        "language": DEFAULT_LANGUAGE,
        # Couleur d'accentuation de l'INTERFACE. Distincte de la couleur des
        # e-mails, qui se règle dans la vue Rédaction et n'est pas enregistrée.
        "ui_accent_color": DEFAULT_ACCENT,
        # Dernière adresse de logo saisie en mode URL, pour la repropose au
        # prochain démarrage. Vide tant que l'utilisateur n'en a pas saisi.
        "logo_url": "",
    }


def read_json(path: Path, fallback: Any) -> Any:
    try:
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, json.JSONDecodeError):
        return fallback


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    with tmp.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    tmp.replace(path)


def ensure_directories() -> None:
    """Crée les dossiers et fichiers de données manquants au démarrage."""
    for directory in (DATA_DIR, UPLOADS_DIR, STATIC_DIR / "css", STATIC_DIR / "js"):
        directory.mkdir(parents=True, exist_ok=True)
    if not CONTACTS_FILE.exists():
        write_json(CONTACTS_FILE, [])
    if not CONFIG_FILE.exists():
        write_json(CONFIG_FILE, default_config())


def load_contacts() -> list[dict[str, Any]]:
    contacts = read_json(CONTACTS_FILE, [])
    return contacts if isinstance(contacts, list) else []


def save_contacts(contacts: list[dict[str, Any]]) -> None:
    write_json(CONTACTS_FILE, contacts)


def load_config() -> dict[str, Any]:
    """Configuration SMTP normalisée (mot de passe encore encodé en base64)."""
    stored = read_json(CONFIG_FILE, {})
    config = default_config()
    if isinstance(stored, dict):
        for key in config:
            if key in stored and stored[key] is not None:
                config[key] = stored[key]
    try:
        config["smtp_port"] = int(config["smtp_port"] or 587)
    except (TypeError, ValueError):
        config["smtp_port"] = 587
    for key in ("smtp_server", "email", "display_name", "password"):
        config[key] = str(config[key] or "").strip()
    config["language"] = normalize_language(config.get("language"))
    config["ui_accent_color"] = normalize_accent(config.get("ui_accent_color"))
    config["logo_url"] = normalize_logo_url(config.get("logo_url")) or ""
    return config


def request_language(request: Request) -> str:
    """Langue d'une requête : en-tête X-Language, sinon celle enregistrée en config.

    L'en-tête permet à l'interface d'appliquer un changement de langue immédiatement,
    y compris avant que le choix ne soit enregistré sur le disque.
    """
    header = request.headers.get(LANGUAGE_HEADER, "")
    if str(header or "").strip():
        return normalize_language(header)
    return load_config()["language"]


def encode_password(plain: str) -> str:
    return base64.b64encode(plain.encode("utf-8")).decode("ascii")


def decode_password(encoded: str) -> str:
    if not encoded:
        return ""
    try:
        return base64.b64decode(encoded.encode("ascii"), validate=True).decode("utf-8")
    except Exception:
        # Valeur saisie à la main en clair dans config.json : acceptée telle quelle.
        return encoded


# --------------------------------------------------------------------------- #
# Utilitaires fichiers, texte & adresses
# --------------------------------------------------------------------------- #


def safe_filename(filename: str) -> str:
    """Nettoie un nom de fichier : pas de chemin, ASCII, caractères sûrs."""
    name = Path(filename or "fichier").name
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._-")
    return name or "fichier"


def unique_upload_path(filename: str) -> Path:
    stem, suffix = os.path.splitext(safe_filename(filename))
    return UPLOADS_DIR / f"{stem[:60] or 'fichier'}-{uuid.uuid4().hex[:8]}{suffix.lower()}"


def upload_basename(value: str) -> str:
    """Accepte « logo.png » ou « /static/uploads/logo.png » et renvoie le nom seul."""
    return Path(str(value or "").split("?")[0].replace("\\", "/")).name


def resolve_upload(filename: str, lang: str = DEFAULT_LANGUAGE) -> Path:
    """Résout un fichier de static/uploads en bloquant toute sortie du dossier."""
    uploads_root = UPLOADS_DIR.resolve()
    candidate = (uploads_root / upload_basename(filename)).resolve()
    if candidate.parent != uploads_root or not candidate.is_file():
        raise HTTPException(status_code=404, detail=tr(lang, "file_not_found", name=filename))
    return candidate


def pretty_attachment_name(stored_name: str) -> str:
    """Retire le suffixe unique ajouté au téléversement (« note-a1b2c3d4.pdf » → « note.pdf »)."""
    stem, suffix = os.path.splitext(Path(stored_name).name)
    return re.sub(r"-[0-9a-f]{8}$", "", stem) + suffix


def display_filename(filename: str) -> str:
    """Nom affiché au destinataire : accents et espaces conservés, chemins exclus.

    Contrairement à safe_filename() (qui protège le disque), on garde l'Unicode :
    Python encode correctement l'en-tête MIME (RFC 2231). Seuls les caractères de
    contrôle sont retirés, pour empêcher toute injection d'en-tête.
    """
    name = Path(str(filename or "").replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:120]


def resolve_attachments(
    attachments: list[Any], lang: str = DEFAULT_LANGUAGE
) -> list[tuple[Path, str]]:
    """Normalise la liste des pièces jointes en couples (chemin, nom vu par le destinataire).

    Chaque entrée est soit un nom de fichier de static/uploads, soit un objet
    ``{"filename": ..., "original_name": ...}`` pour conserver le nom d'origine.
    """
    resolved: list[tuple[Path, str]] = []
    for item in attachments or []:
        if isinstance(item, dict):
            stored = str(item.get("filename") or "")
            label = display_filename(str(item.get("original_name") or "")) or None
        else:
            stored = str(item or "")
            label = None
        if not stored.strip():
            continue
        path = resolve_upload(stored, lang)
        resolved.append((path, label or pretty_attachment_name(path.name)))
    return resolved


def is_valid_email(value: str) -> bool:
    return bool(EMAIL_RE.match((value or "").strip()))


def split_recipients(values: Any) -> list[str]:
    """Accepte une liste ou « a@b.com, c@d.com » et renvoie des adresses dédoublonnées."""
    if not values:
        return []
    if isinstance(values, str):
        values = [values]
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        for part in re.split(r"[,;\s]+", str(value or "")):
            addr = part.strip().strip("<>")
            if addr and addr.lower() not in seen:
                seen.add(addr.lower())
                out.append(addr)
    return out


def text_to_html(text: str) -> Markup:
    """Échappe le texte saisi puis convertit les sauts de ligne en paragraphes/<br>."""
    escaped = str(escape(text or ""))
    paragraphs = [p for p in re.split(r"\n\s*\n", escaped) if p.strip()]
    blocks = [p.strip("\n").replace("\n", "<br>") for p in paragraphs]
    return Markup("".join(f'<p style="margin:0 0 16px 0;">{b}</p>' for b in blocks))


def signature_to_html(text: str) -> Markup:
    """Signature : une ligne = un <br>, la première ligne en gras."""
    lines = [str(escape(line)) for line in (text or "").splitlines() if line.strip()]
    if not lines:
        return Markup("")
    lines[0] = f"<strong>{lines[0]}</strong>"
    return Markup("<br>".join(lines))


def to_plain_text(
    title: str,
    body: str,
    signature: str,
    company: str,
    extras: list[str] | None = None,
) -> str:
    """Version texte brut de secours (multipart/alternative).

    ``extras`` porte les liens des blocs optionnels : sans eux, un destinataire
    lisant la version texte ne verrait ni le bouton d'action ni la video.
    """
    parts = [p.strip() for p in (title, body, signature, company) if p and p.strip()]
    parts += [p.strip() for p in (extras or []) if p and p.strip()]
    text = "\n\n".join(parts)
    text = re.sub(r"<[^>]+>", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def normalize_web_url(value: Any) -> str | None:
    """Valide une adresse publique destinee a un e-mail.

    Renvoie l'URL nettoyée, "" si le champ est vide, ou None si l'adresse n'est
    pas exploitable. Seuls http et https sont acceptés : liens et images d'un
    e-mail sont ouverts par le client de messagerie du destinataire, où tout
    autre schéma (``javascript:``, ``file:``) serait au mieux inerte.
    """
    url = str(value or "").strip()
    if not url:
        return ""
    if not re.match(r"^https?://[^\s]+$", url, re.IGNORECASE):
        return None
    return url[:2000]


def normalize_logo_url(value: Any) -> str | None:
    """Valide l'adresse publique d'un logo. Voir normalize_web_url()."""
    return normalize_web_url(value)


def normalize_accent(color: Any) -> str:
    color = str(color or "").strip()
    return color if re.fullmatch(r"#[0-9A-Fa-f]{6}", color) else DEFAULT_ACCENT


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


# --------------------------------------------------------------------------- #
# Modèles d'e-mail (Jinja2)
# --------------------------------------------------------------------------- #

email_env = Environment(
    loader=FileSystemLoader(str(EMAIL_TEMPLATES_DIR)),
    autoescape=select_autoescape(["html", "xml"]),
)


def available_templates() -> list[str]:
    """Noms des modèles présents dans templates/email_templates (sans extension)."""
    if not EMAIL_TEMPLATES_DIR.is_dir():
        return []
    names = sorted(p.stem for p in EMAIL_TEMPLATES_DIR.glob("*.html") if p.is_file())
    order = list(TEMPLATE_ORDER)
    return sorted(names, key=lambda n: (order.index(n) if n in order else len(order), n))


# --------------------------------------------------------------------------- #
# Blocs optionnels injectes dans les modeles
#
# Les trois blocs sont construits ici, en Python, puis passes aux modeles sous
# forme de Markup. Un modele se contente de placer {{ video_block }},
# {{ cta_block }} et {{ social_block }} : le balisage est ecrit une seule fois
# pour les 10 modeles, la ou le dupliquer dix fois garantirait de les voir
# diverger. Un champ vide produit une chaine vide, donc aucune ligne de tableau
# et aucun blanc dans le message.
# --------------------------------------------------------------------------- #

# Glyphes simples, volontairement geometriques et monochromes : une seule
# couleur a substituer, et aucun fichier externe a heberger.
SOCIAL_NETWORKS: tuple[tuple[str, str, str], ...] = (
    ("instagram", "Instagram",
     '<rect x="3" y="3" width="18" height="18" rx="5" fill="none" stroke="{c}"'
     ' stroke-width="2"/><circle cx="12" cy="12" r="4" fill="none" stroke="{c}"'
     ' stroke-width="2"/><circle cx="17.2" cy="6.8" r="1.3" fill="{c}"/>'),
    ("facebook", "Facebook",
     '<path d="M13.6 21v-8h2.6l.4-3h-3V8.1c0-.87.25-1.46 1.5-1.46h1.6V4.06'
     'c-.28-.04-1.2-.12-2.26-.12-2.2 0-3.72 1.35-3.72 3.82V10H8.2v3h2.52v8z"'
     ' fill="{c}"/>'),
    ("youtube", "YouTube",
     '<rect x="2.5" y="6" width="19" height="12" rx="3.5" fill="none"'
     ' stroke="{c}" stroke-width="2"/><path d="M10.6 9.6l4.9 2.4-4.9 2.4z"'
     ' fill="{c}"/>'),
    ("linkedin", "LinkedIn",
     '<rect x="3" y="3" width="18" height="18" rx="3" fill="none" stroke="{c}"'
     ' stroke-width="2"/><path d="M7.3 10.4v6.2" stroke="{c}" stroke-width="2"'
     ' stroke-linecap="round"/><circle cx="7.3" cy="7.6" r="1.1" fill="{c}"/>'
     '<path d="M11.3 16.6v-3.4a2.3 2.3 0 0 1 4.6 0v3.4" fill="none"'
     ' stroke="{c}" stroke-width="2" stroke-linecap="round"/>'),
    ("x", "X",
     '<path d="M4.6 4.6l14.8 14.8M19.4 4.6L4.6 19.4" stroke="{c}"'
     ' stroke-width="2.4" stroke-linecap="round"/>'),
    ("tiktok", "TikTok",
     '<path d="M14.2 4.2v9.1a3.3 3.3 0 1 1-2.7-3.24" fill="none" stroke="{c}"'
     ' stroke-width="2" stroke-linecap="round" stroke-linejoin="round"/>'
     '<path d="M14.2 4.2c.5 2.2 2 3.6 4.3 3.9" fill="none" stroke="{c}"'
     ' stroke-width="2" stroke-linecap="round"/>'),
)

FONT_STACK = "'Helvetica Neue',Helvetica,Arial,sans-serif"


def readable_on(background: str) -> str:
    """Noir ou blanc, selon ce qui se lit le mieux sur ``background``.

    Meme calcul de luminance que applyUiAccent() cote interface : une couleur
    d'accentuation claire recoit un glyphe sombre, et inversement.
    """
    color = normalize_accent(background).lstrip("#")
    r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "#0f172a" if luminance > 0.68 else "#ffffff"


def social_links(variables: dict[str, Any]) -> list[tuple[str, str, str, str]]:
    """Reseaux renseignes, dans l'ordre de SOCIAL_NETWORKS. Adresses validees.

    Renvoie des quadruplets (cle, libelle, adresse, glyphe) : la cle sert a
    retrouver l'icone hebergee correspondante dans SOCIAL_ICON_URLS.
    """
    out: list[tuple[str, str, str, str]] = []
    for key, label, glyph in SOCIAL_NETWORKS:
        url = normalize_web_url(variables.get(f"social_{key}"))
        if url:
            out.append((key, label, url, glyph))
    return out


def social_badge(key: str, label: str, glyph: str, accent: str) -> str:
    """Contenu d'une pastille : image hebergee si elle est configuree, sinon SVG.

    Une image hebergee (constants.SOCIAL_ICON_URLS) est le seul rendu qui tienne
    dans un vrai client de messagerie : Gmail et Outlook suppriment les <svg>, et
    integrer les six PNG au message ferait apparaitre six pieces jointes. Le PNG
    porte deja sa pastille ronde et sa couleur de marque, la cellule n'en ajoute
    donc aucune. Sans adresse configuree, on retombe sur le SVG dans une pastille
    a la couleur de l'e-mail : l'apercu reste juste, et aucune image cassee n'est
    envoyee.
    """
    icon_url = normalize_web_url(SOCIAL_ICON_URLS.get(key)) or ""
    size = SOCIAL_ICON_SIZE
    if icon_url:
        inner = (
            '<img src="%s" width="%d" height="%d" alt="%s"'
            ' style="display:block; width:%dpx; height:%dpx; border:0;'
            ' outline:none; text-decoration:none; border-radius:%dpx;">'
            % (escape(icon_url), size, size, escape(label), size, size, size // 2)
        )
        cell_style = ""
    else:
        inner = (
            '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"'
            ' viewBox="0 0 24 24" role="img" aria-label="%s">%s</svg>'
            % (escape(label), glyph.format(c=readable_on(accent)))
        )
        cell_style = ' bgcolor="%s"' % accent
    return (
        '<td align="center" valign="middle" width="%d" height="%d"%s'
        ' style="width:%dpx; height:%dpx; border-radius:%dpx; line-height:0;%s">'
        '<a href="%%s" target="_blank" title="%%s"'
        ' style="display:inline-block; line-height:0; text-decoration:none;">'
        "%s</a></td>"
        '<td width="8" style="width:8px; line-height:0; font-size:0;">&nbsp;</td>'
        % (size, size, cell_style, size, size, size // 2,
           (" background-color:%s;" % accent) if not icon_url else "", inner)
    )


def build_social_block(variables: dict[str, Any], accent: str) -> Markup:
    """Rangee de pastilles rondes cliquables, pour le pied de page."""
    links = social_links(variables)
    if not links:
        return Markup("")
    cells = [
        social_badge(key, label, glyph, accent) % (escape(url), escape(label))
        for key, label, url, glyph in links
    ]
    return Markup(
        '<table role="presentation" cellpadding="0" cellspacing="0" border="0"'
        ' align="center" style="margin:0 auto 14px auto;"><tr>%s</tr></table>'
        % "".join(cells)
    )


def build_cta_block(variables: dict[str, Any], accent: str) -> Markup:
    """Bouton d'action, en tableau avec repli VML pour Outlook.

    Outlook (moteur Word) ignore padding et border-radius sur un lien : sans le
    v:roundrect, le bouton s'y reduirait a du texte souligne.
    """
    text = str(variables.get("cta_text") or "").strip()
    url = normalize_web_url(variables.get("cta_url"))
    if not text or not url:
        return Markup("")
    label = escape(text[:80])
    href = escape(url)
    fg = readable_on(accent)
    return Markup(
        '<tr><td align="center" style="padding:8px 40px 28px 40px;">'
        '<table role="presentation" cellpadding="0" cellspacing="0" border="0"'
        ' align="center"><tr><td align="center" bgcolor="%(accent)s"'
        ' style="background-color:%(accent)s; border-radius:8px;">'
        "<!--[if mso]>"
        '<v:roundrect xmlns:v="urn:schemas-microsoft-com:vml"'
        ' xmlns:w="urn:schemas-microsoft-com:office:word" href="%(href)s"'
        ' style="height:48px; v-text-anchor:middle; width:280px;" arcsize="17%%"'
        ' strokecolor="%(accent)s" fillcolor="%(accent)s"><w:anchorlock/>'
        '<center style="color:%(fg)s; font-family:%(font)s; font-size:16px;'
        ' font-weight:bold;">%(label)s</center></v:roundrect>'
        "<![endif]-->"
        "<!--[if !mso]><!-->"
        '<a href="%(href)s" target="_blank"'
        ' style="display:inline-block; padding:15px 34px; font-family:%(font)s;'
        ' font-size:16px; line-height:18px; font-weight:700; color:%(fg)s;'
        ' text-decoration:none; border-radius:8px; background-color:%(accent)s;">'
        "%(label)s</a>"
        "<!--<![endif]-->"
        "</td></tr></table></td></tr>"
        % {"accent": accent, "href": href, "fg": fg, "label": label,
           "font": FONT_STACK}
    )


# Caractere de lecture : U+25BA, « black right-pointing pointer ». Choisi
# plutot que U+25B6 parce qu'il n'a pas de variante emoji : aucun client ne le
# transformera en pictogramme colore. Un <svg> aurait ete supprime par Gmail et
# Outlook, et une image hebergee ne pourrait pas suivre accent_color.
PLAY_GLYPH = "\u25ba"


def build_video_block(
    variables: dict[str, Any], accent: str, lang: str = DEFAULT_LANGUAGE
) -> Markup:
    """Vignette cliquable avec pastille de lecture par-dessus.

    L'image est posee en fond de cellule (avec repli VML pour Outlook) afin que
    la pastille se superpose vraiment : un e-mail ne peut pas empiler deux
    elements en CSS. Le clic ouvre la video dans le navigateur, aucun client de
    messagerie ne sachant lire une video en ligne.

    Le triangle est un caractere, pas un dessin : il s'affiche partout, y compris
    la ou un <svg> serait supprime, et prend la couleur de l'e-mail.
    """
    thumb = normalize_web_url(variables.get("video_thumbnail"))
    url = normalize_web_url(variables.get("video_url"))
    if not thumb or not url:
        return Markup("")
    src = escape(thumb)
    href = escape(url)
    label = escape(email_strings(lang)["video_link"])
    play = (
        '<table role="presentation" cellpadding="0" cellspacing="0" border="0"'
        ' align="center"><tr>'
        '<td align="center" valign="middle" width="64" height="64" bgcolor="#ffffff"'
        ' style="width:64px; height:64px; background-color:rgba(255,255,255,0.92);'
        ' border-radius:32px; font-family:%s; font-size:27px; line-height:64px;'
        ' mso-line-height-rule:exactly; text-indent:4px; color:%s;">%s'
        "</td></tr></table>" % (FONT_STACK, accent, PLAY_GLYPH)
    )
    return Markup(
        '<tr><td align="center" style="padding:4px 40px 20px 40px;">'
        '<a href="%(href)s" target="_blank" title="%(label)s"'
        ' aria-label="%(label)s" style="text-decoration:none;">'
        '<table role="presentation" width="100%%" cellpadding="0" cellspacing="0"'
        ' border="0"><tr>'
        '<td background="%(src)s" bgcolor="#0f172a" height="260" valign="middle"'
        ' align="center" style="height:260px; background-color:#0f172a;'
        " background-image:url('%(src)s'); background-position:center center;"
        ' background-size:cover; background-repeat:no-repeat; border-radius:10px;">'
        "<!--[if gte mso 9]>"
        '<v:rect xmlns:v="urn:schemas-microsoft-com:vml" fill="true"'
        ' stroke="false" style="width:520px; height:260px;">'
        '<v:fill type="frame" src="%(src)s" color="#0f172a"/>'
        '<v:textbox inset="0,0,0,0">'
        "<![endif]-->"
        "%(play)s"
        "<!--[if gte mso 9]></v:textbox></v:rect><![endif]-->"
        "</td></tr></table></a></td></tr>"
        % {"href": href, "src": src, "play": play, "label": label}
    )


def plain_text_extras(variables: dict[str, Any], lang: str) -> list[str]:
    """Liens des blocs optionnels, pour la version texte du message."""
    strings = email_strings(lang)
    out: list[str] = []
    cta_text = str(variables.get("cta_text") or "").strip()
    cta_url = normalize_web_url(variables.get("cta_url"))
    if cta_text and cta_url:
        out.append(f"{cta_text} : {cta_url}")
    video_url = normalize_web_url(variables.get("video_url"))
    if normalize_web_url(variables.get("video_thumbnail")) and video_url:
        out.append(f"{strings['video_link']} : {video_url}")
    links = social_links(variables)
    if links:
        out.append(
            strings["social_follow"]
            + " "
            + " | ".join(f"{label} : {url}" for _, label, url, _ in links)
        )
    return out


def render_email_html(
    template: str,
    variables: dict[str, Any],
    logo_src: str = "",
    lang: str = DEFAULT_LANGUAGE,
) -> str:
    """Rend un modèle d'e-mail avec les variables de l'utilisateur.

    ``lang`` ne traduit que les textes fixes du modèle (pied de page, mentions
    légales…) : le titre, le corps et la signature restent exactement ce que
    l'utilisateur a saisi, quelle que soit la langue de l'interface.
    """
    name = (template or "newsletter").strip()
    code = normalize_language(lang)
    if name not in available_templates():
        raise HTTPException(status_code=404, detail=tr(code, "unknown_template", name=template))

    title = str(variables.get("title") or "").strip()
    company = str(variables.get("company_name") or "").strip()
    accent = normalize_accent(variables.get("accent_color"))
    context = {
        "logo_url": logo_src,
        "title": title,
        "body": text_to_html(str(variables.get("body") or "")),
        "body_text": str(variables.get("body") or ""),
        "signature": signature_to_html(str(variables.get("signature") or "")),
        "signature_text": str(variables.get("signature") or ""),
        "company_name": company,
        "accent_color": accent,
        "social_block": build_social_block(variables, accent),
        "cta_block": build_cta_block(variables, accent),
        "video_block": build_video_block(variables, accent, code),
        "year": datetime.now().year,
        "lang": code,
        "t": email_strings(code),
    }
    try:
        return email_env.get_template(f"{name}.html").render(**context)
    except TemplateNotFound as exc:  # pragma: no cover - garde-fou
        raise HTTPException(
            status_code=404, detail=tr(code, "unknown_template", name=template)
        ) from exc


def logo_preview_src(logo: Any) -> str:
    """URL utilisable dans l'aperçu navigateur (URL publique ou fichier uploadé)."""
    value = str(logo or "").strip()
    if not value:
        return ""
    if value.startswith(("http://", "https://", "data:", "/static/")):
        return value
    return f"/static/uploads/{upload_basename(value)}"


def logo_data_uri(path: Path) -> str:
    """Encode une image locale en data URI base64 (utilisé pour l'aperçu autonome)."""
    mime = mimetypes.guess_type(path.name)[0] or "image/png"
    payload = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{payload}"


# --------------------------------------------------------------------------- #
# SMTP
# --------------------------------------------------------------------------- #


class SmtpError(Exception):
    """Erreur SMTP dont le message est déjà traduit dans la langue de la requête."""


def smtp_error_message(
    exc: Exception, server: str = "", port: int = 0, lang: str = DEFAULT_LANGUAGE
) -> str:
    """Traduit une exception SMTP en message lisible dans la langue de l'interface."""
    where = {"server": server, "port": port}
    if isinstance(exc, smtplib.SMTPAuthenticationError):
        return tr(lang, "smtp_auth")
    if isinstance(exc, smtplib.SMTPRecipientsRefused):
        return tr(lang, "smtp_recipient_refused")
    if isinstance(exc, smtplib.SMTPSenderRefused):
        return tr(lang, "smtp_sender_refused")
    if isinstance(exc, smtplib.SMTPDataError):
        return tr(lang, "smtp_data", error=exc)
    if isinstance(exc, smtplib.SMTPConnectError):
        return tr(lang, "smtp_connect", **where)
    if isinstance(exc, smtplib.SMTPServerDisconnected):
        return tr(lang, "smtp_disconnected")
    if isinstance(exc, ssl.SSLError):
        return tr(lang, "smtp_ssl")
    if isinstance(exc, TimeoutError):
        return tr(lang, "smtp_timeout", **where)
    if isinstance(exc, OSError):
        return tr(lang, "smtp_network", error=exc.strerror or exc, **where)
    if isinstance(exc, smtplib.SMTPException):
        return tr(lang, "smtp_generic", error=exc)
    return tr(lang, "smtp_unexpected", error=exc)


def validate_smtp_config(
    config: dict[str, Any], lang: str = DEFAULT_LANGUAGE
) -> tuple[str, int, str, str, str]:
    """Vérifie la configuration et renvoie (serveur, port, email, mot de passe, nom)."""
    server = str(config.get("smtp_server") or "").strip()
    email_addr = str(config.get("email") or "").strip()
    password = decode_password(str(config.get("password") or ""))
    display_name = str(config.get("display_name") or "").strip()
    try:
        port = int(config.get("smtp_port") or 587)
    except (TypeError, ValueError):
        port = 587

    missing = []
    if not server:
        missing.append(tr(lang, "missing_server"))
    if not email_addr:
        missing.append(tr(lang, "missing_email"))
    if not password:
        missing.append(tr(lang, "missing_password"))
    if missing:
        raise HTTPException(
            status_code=400,
            detail=tr(lang, "config_incomplete", missing=", ".join(missing)),
        )
    if not is_valid_email(email_addr):
        raise HTTPException(
            status_code=400,
            detail=tr(lang, "invalid_sender", email=email_addr),
        )
    return server, port, email_addr, password, display_name


class SmtpSession:
    """Connexion SMTP réutilisable pour l'envoi d'une campagne (code bloquant)."""

    def __init__(
        self,
        server: str,
        port: int,
        email_addr: str,
        password: str,
        lang: str = DEFAULT_LANGUAGE,
    ) -> None:
        self.server = server
        self.port = port
        self.email = email_addr
        self.password = password
        self.lang = lang
        self.client: smtplib.SMTP | None = None

    def connect(self) -> None:
        context = ssl.create_default_context()
        try:
            if self.port == 465:
                client: smtplib.SMTP = smtplib.SMTP_SSL(
                    self.server, self.port, timeout=30, context=context
                )
            else:
                client = smtplib.SMTP(self.server, self.port, timeout=30)
                client.ehlo()
                if client.has_extn("starttls"):
                    client.starttls(context=context)
                    client.ehlo()
            client.login(self.email, self.password)
        except Exception as exc:
            raise SmtpError(smtp_error_message(exc, self.server, self.port, self.lang)) from exc
        self.client = client

    def send(self, message: EmailMessage) -> None:
        if self.client is None:
            self.connect()
        assert self.client is not None
        try:
            self.client.send_message(message)
        except (smtplib.SMTPServerDisconnected, smtplib.SMTPConnectError):
            # Connexion perdue entre deux envois : on rouvre puis on réessaie une fois.
            self.close()
            self.connect()
            assert self.client is not None
            try:
                self.client.send_message(message)
            except Exception as exc:
                raise SmtpError(smtp_error_message(exc, self.server, self.port, self.lang)) from exc
        except Exception as exc:
            raise SmtpError(smtp_error_message(exc, self.server, self.port, self.lang)) from exc

    def close(self) -> None:
        if self.client is not None:
            try:
                self.client.quit()
            except Exception:
                pass
            self.client = None

    def __enter__(self) -> SmtpSession:
        self.connect()
        return self

    def __exit__(self, *_exc_info: Any) -> None:
        self.close()


def build_message(
    *,
    sender: str,
    display_name: str,
    recipient: str,
    subject: str,
    html: str,
    text: str,
    logo_cid: str | None,
    logo_file: Path | None,
    attachments: list[tuple[Path, str]],
) -> EmailMessage:
    """Construit un e-mail multipart (texte + HTML, logo inline, pièces jointes)."""
    message = EmailMessage()
    message["From"] = formataddr((display_name or sender, sender))
    message["To"] = recipient
    message["Subject"] = subject
    message["Date"] = formatdate(localtime=True)
    message["Message-ID"] = make_msgid(domain=sender.split("@")[-1] or "localhost")

    message.set_content(text or " ")
    message.add_alternative(html, subtype="html")

    if logo_file is not None and logo_cid:
        mime = mimetypes.guess_type(logo_file.name)[0] or "image/png"
        maintype, _, subtype = mime.partition("/")
        html_part = message.get_payload()[-1]
        html_part.add_related(
            logo_file.read_bytes(),
            maintype=maintype or "image",
            subtype=subtype or "png",
            cid=f"<{logo_cid}>",
            filename=logo_file.name,
            disposition="inline",
        )

    for path, display_name in attachments:
        mime = mimetypes.guess_type(display_name or path.name)[0] or "application/octet-stream"
        maintype, _, subtype = mime.partition("/")
        message.add_attachment(
            path.read_bytes(),
            maintype=maintype or "application",
            subtype=subtype or "octet-stream",
            filename=display_name or path.name,
        )
    return message


def send_campaign(
    *,
    config: dict[str, Any],
    recipients: list[str],
    subject: str,
    template: str,
    variables: dict[str, Any],
    attachments: list[Any],
    delay: float,
    lang: str = DEFAULT_LANGUAGE,
) -> dict[str, Any]:
    """Envoie un e-mail par destinataire (code bloquant, lancé dans un thread)."""
    server, port, sender, password, display_name = validate_smtp_config(config, lang)
    attachment_paths = resolve_attachments(attachments, lang)

    # Logo : un fichier uploadé est intégré au message (cid), une URL est utilisée telle quelle.
    raw_logo = str(variables.get("logo_url") or "").strip()
    logo_file: Path | None = None
    logo_cid: str | None = None
    logo_src = ""
    if raw_logo:
        if raw_logo.startswith(("http://", "https://", "data:")):
            logo_src = raw_logo
        else:
            logo_file = resolve_upload(raw_logo, lang)
            logo_cid = make_msgid(domain="emailsender.local")[1:-1]
            logo_src = f"cid:{logo_cid}"

    html = render_email_html(template, variables, logo_src=logo_src, lang=lang)
    text = to_plain_text(
        str(variables.get("title") or ""),
        str(variables.get("body") or ""),
        str(variables.get("signature") or ""),
        str(variables.get("company_name") or ""),
        plain_text_extras(variables, lang),
    )

    results: list[dict[str, Any]] = []
    session = SmtpSession(server, port, sender, password, lang)
    try:
        try:
            session.connect()
        except SmtpError as exc:
            # Échec avant le premier envoi : rien n'est parti, on remonte l'erreur telle quelle.
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        for index, recipient in enumerate(recipients):
            if index and delay > 0:
                time.sleep(delay)
            try:
                message = build_message(
                    sender=sender,
                    display_name=display_name,
                    recipient=recipient,
                    subject=subject,
                    html=html,
                    text=text,
                    logo_cid=logo_cid,
                    logo_file=logo_file,
                    attachments=attachment_paths,
                )
                session.send(message)
                results.append(
                    {"email": recipient, "ok": True, "message": tr(lang, "result_sent")}
                )
            except SmtpError as exc:
                results.append({"email": recipient, "ok": False, "message": str(exc)})
            except Exception as exc:  # pragma: no cover - garde-fou
                results.append(
                    {
                        "email": recipient,
                        "ok": False,
                        "message": tr(lang, "result_error", error=exc),
                    }
                )
    finally:
        session.close()

    sent = sum(1 for r in results if r["ok"])
    failed = len(results) - sent
    if failed == 0:
        summary = tr(lang, "summary_all_sent", sent=sent)
    elif sent == 0:
        summary = tr(lang, "summary_none_sent", failed=failed)
    else:
        summary = tr(lang, "summary_mixed", sent=sent, failed=failed)
    return {"ok": failed == 0, "sent": sent, "failed": failed, "results": results,
            "message": summary}


def test_smtp_connection(
    config: dict[str, Any], lang: str = DEFAULT_LANGUAGE
) -> dict[str, Any]:
    """Tente une connexion + authentification SMTP (code bloquant)."""
    server, port, sender, password, _ = validate_smtp_config(config, lang)
    session = SmtpSession(server, port, sender, password, lang)
    try:
        session.connect()
    except SmtpError as exc:
        return {"ok": False, "message": str(exc)}
    finally:
        session.close()
    return {
        "ok": True,
        "message": tr(lang, "test_success", server=server, port=port, email=sender),
    }


# --------------------------------------------------------------------------- #
# Schémas Pydantic
# --------------------------------------------------------------------------- #


class EmailVariables(BaseModel):
    logo_url: str = ""
    title: str = ""
    body: str = ""
    signature: str = ""
    company_name: str = ""
    accent_color: str = DEFAULT_ACCENT
    # Blocs optionnels. Une adresse non conforme est ignoree au rendu plutot
    # que refusee : un champ mal saisi ne doit pas bloquer tout un envoi.
    social_instagram: str = ""
    social_facebook: str = ""
    social_youtube: str = ""
    social_linkedin: str = ""
    social_x: str = ""
    social_tiktok: str = ""
    cta_text: str = ""
    cta_url: str = ""
    video_thumbnail: str = ""
    video_url: str = ""


class SendRequest(BaseModel):
    to: list[str] | str = Field(default_factory=list)
    subject: str = ""
    template: str = "newsletter"
    variables: EmailVariables = Field(default_factory=EmailVariables)
    # Nom de fichier simple, ou { filename, original_name } pour garder le nom d'origine.
    attachments: list[str | dict[str, str]] = Field(default_factory=list)
    delay: float | None = None


class PreviewRequest(BaseModel):
    template: str = "newsletter"
    variables: EmailVariables = Field(default_factory=EmailVariables)


class ContactIn(BaseModel):
    name: str = ""
    email: str = ""
    group: str = ""


class ContactsIn(BaseModel):
    """Accepte un contact seul ou une liste via le champ « contacts »."""

    name: str = ""
    email: str = ""
    group: str = ""
    contacts: list[ContactIn] | None = None


class ConfigIn(BaseModel):
    smtp_server: str = ""
    smtp_port: int | str = 587
    email: str = ""
    password: str = ""
    display_name: str = ""
    # Optionnels : le formulaire SMTP ne modifie ni la langue (/api/language)
    # ni la couleur de l'interface (/api/theme).
    language: str = ""
    ui_accent_color: str = ""


class LanguageIn(BaseModel):
    language: str = DEFAULT_LANGUAGE


class ThemeIn(BaseModel):
    ui_accent_color: str = DEFAULT_ACCENT


class LogoIn(BaseModel):
    logo_url: str = ""


# --------------------------------------------------------------------------- #
# Application
# --------------------------------------------------------------------------- #

ensure_directories()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    ensure_directories()
    yield


class RevalidatingStaticFiles(StaticFiles):
    """Sert /static en imposant une revalidation avant réutilisation.

    Sans en-tête Cache-Control, les navigateurs appliquent un cache heuristique et
    peuvent resservir un fichier périmé sans consulter le serveur, laissant tourner
    un app.js obsolète après une mise à jour. « no-cache » n'interdit pas la mise en
    cache, il impose la revalidation ; l'ETag la rend gratuite (304, aucun octet
    retransmis).
    """

    def file_response(self, *args: Any, **kwargs: Any) -> Any:
        response = super().file_response(*args, **kwargs)
        response.headers["Cache-Control"] = "no-cache, must-revalidate"
        return response


app = FastAPI(title="Zm@il", version="1.2.0", lifespan=lifespan)
app.mount("/static", RevalidatingStaticFiles(directory=str(STATIC_DIR)), name="static")
ui_templates = Jinja2Templates(directory=str(TEMPLATES_DIR))


def asset_version() -> str:
    """Empreinte des feuilles de style et scripts, ajoutée en ?v= à leur URL.

    Change dès qu'un fichier est modifié : le navigateur voit alors une URL
    inédite et ne peut pas servir une version périmée.
    """
    stamp = 0.0
    for name in ("css/style.css", "js/i18n.js", "js/app.js"):
        path = STATIC_DIR / name
        try:
            stamp = max(stamp, path.stat().st_mtime)
        except OSError:
            continue
    return str(int(stamp))


@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException) -> JSONResponse:
    """Toutes les erreurs API répondent avec { ok: false, message }."""
    lang = request_language(request)
    detail = exc.detail if isinstance(exc.detail, str) else tr(lang, "bad_request")
    return JSONResponse(status_code=exc.status_code, content={"ok": False, "message": detail})


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(
    request: Request, exc: RequestValidationError
) -> JSONResponse:
    """Erreurs de validation traduites en un message lisible."""
    lang = request_language(request)
    first = exc.errors()[0] if exc.errors() else {}
    field = ".".join(str(part) for part in first.get("loc", ())[1:]) or tr(lang, "request_field")
    return JSONResponse(
        status_code=422,
        content={"ok": False, "message": tr(lang, "invalid_data", field=field)},
    )


# -------------------------------- Interface -------------------------------- #


@app.get("/", response_class=HTMLResponse)
async def index(request: Request) -> HTMLResponse:
    """Sert l'interface déjà réglée (langue et couleur) : aucun clignotement."""
    config = load_config()
    language = config["language"]
    response = ui_templates.TemplateResponse(
        request,
        "index.html",
        {
            "default_accent": DEFAULT_ACCENT,
            "year": datetime.now().year,
            "asset_version": asset_version(),
            "language": language,
            "ui_accent_color": config["ui_accent_color"],
            "logo_url": config["logo_url"],
            "languages": [
                {"code": code, "label": LANGUAGE_LABELS[code]} for code in LANGUAGES
            ],
        },
    )
    response.headers["Cache-Control"] = "no-store"
    return response


# ---------------------------------- Envoi ---------------------------------- #


def _prepare_send(
    payload: SendRequest, lang: str = DEFAULT_LANGUAGE
) -> tuple[list[str], str, dict[str, Any]]:
    recipients = split_recipients(payload.to)
    if not recipients:
        raise HTTPException(status_code=400, detail=tr(lang, "no_recipients"))

    invalid = [addr for addr in recipients if not is_valid_email(addr)]
    if invalid:
        raise HTTPException(
            status_code=400,
            detail=tr(lang, "invalid_addresses", addresses=", ".join(invalid[:5])),
        )

    subject = (payload.subject or "").strip()
    if not subject:
        raise HTTPException(status_code=400, detail=tr(lang, "subject_required"))

    variables = payload.variables.model_dump()
    if not str(variables.get("body") or "").strip() and not str(variables.get("title") or "").strip():
        raise HTTPException(status_code=400, detail=tr(lang, "empty_message"))
    return recipients, subject, variables


@app.post("/api/send")
async def api_send(
    payload: SendRequest, lang: str = Depends(request_language)
) -> dict[str, Any]:
    """Envoi vers un ou plusieurs destinataires."""
    recipients, subject, variables = _prepare_send(payload, lang)
    config = load_config()
    delay = payload.delay if payload.delay is not None else (
        BULK_DELAY_SECONDS if len(recipients) > 1 else 0.0
    )
    return await asyncio.to_thread(
        lambda: send_campaign(
            config=config,
            recipients=recipients,
            subject=subject,
            template=payload.template,
            variables=variables,
            attachments=payload.attachments,
            delay=max(0.0, float(delay)),
            lang=lang,
        )
    )


@app.post("/api/send-bulk")
async def api_send_bulk(
    payload: SendRequest, lang: str = Depends(request_language)
) -> dict[str, Any]:
    """Envoi en masse : une pause (1 s par défaut) entre chaque destinataire."""
    recipients, subject, variables = _prepare_send(payload, lang)
    config = load_config()
    delay = BULK_DELAY_SECONDS if payload.delay is None else max(0.0, float(payload.delay))
    return await asyncio.to_thread(
        lambda: send_campaign(
            config=config,
            recipients=recipients,
            subject=subject,
            template=payload.template,
            variables=variables,
            attachments=payload.attachments,
            delay=delay,
            lang=lang,
        )
    )


# ------------------------------ Téléversement ------------------------------ #


@app.post("/api/upload")
async def api_upload(
    file: UploadFile = File(...),
    kind: str = Form("attachment"),
    lang: str = Depends(request_language),
) -> dict[str, Any]:
    """Enregistre un logo ou une pièce jointe dans static/uploads."""
    ensure_directories()
    if not file.filename:
        raise HTTPException(status_code=400, detail=tr(lang, "no_file"))

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail=tr(lang, "empty_file"))
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=tr(
                lang,
                "file_too_large",
                size=len(content) // 1024,
                limit=MAX_UPLOAD_BYTES // (1024 * 1024),
            ),
        )

    mime = file.content_type or mimetypes.guess_type(file.filename)[0] or ""
    if kind == "logo" and not mime.startswith("image/"):
        raise HTTPException(status_code=400, detail=tr(lang, "logo_not_image"))

    destination = unique_upload_path(file.filename)
    destination.write_bytes(content)
    return {
        "ok": True,
        "message": tr(lang, "upload_done", name=file.filename),
        "filename": destination.name,
        "original_name": file.filename,
        "url": f"/static/uploads/{destination.name}",
        "size": len(content),
        "content_type": mime,
        "is_image": mime.startswith("image/"),
    }


# -------------------------------- Contacts --------------------------------- #


def normalized_contact(name: str, email_addr: str, group: str) -> dict[str, Any]:
    return {
        "id": uuid.uuid4().hex[:12],
        "name": (name or "").strip() or (email_addr or "").split("@")[0],
        "email": (email_addr or "").strip(),
        "group": (group or "").strip() or DEFAULT_GROUP,
        "created_at": now_iso(),
    }


def add_contacts(
    rows: list[dict[str, str]], lang: str = DEFAULT_LANGUAGE
) -> dict[str, Any]:
    """Ajoute des contacts en ignorant les doublons et les adresses invalides."""
    contacts = load_contacts()
    known = {str(c.get("email", "")).strip().lower() for c in contacts}
    added: list[dict[str, Any]] = []
    duplicates: list[str] = []
    invalid: list[str] = []

    for row in rows:
        email_addr = str(row.get("email") or "").strip()
        if not is_valid_email(email_addr):
            invalid.append(email_addr or "(vide)")
            continue
        if email_addr.lower() in known:
            duplicates.append(email_addr)
            continue
        contact = normalized_contact(str(row.get("name") or ""), email_addr, str(row.get("group") or ""))
        contacts.append(contact)
        known.add(email_addr.lower())
        added.append(contact)

    if added:
        save_contacts(contacts)

    pieces = [tr(lang, "contacts_added", count=len(added))]
    if duplicates:
        pieces.append(tr(lang, "contacts_duplicates", count=len(duplicates)))
    if invalid:
        pieces.append(tr(lang, "contacts_invalid", count=len(invalid)))
    return {
        "ok": bool(added) or not rows,
        "message": ", ".join(pieces) + ".",
        "added": added,
        "duplicates": duplicates,
        "invalid": invalid,
        "contacts": contacts,
    }


@app.get("/api/contacts")
async def api_contacts(q: str = "", group: str = "") -> dict[str, Any]:
    contacts = load_contacts()
    needle = q.strip().lower()
    if needle:
        contacts = [
            c
            for c in contacts
            if needle in str(c.get("name", "")).lower()
            or needle in str(c.get("email", "")).lower()
            or needle in str(c.get("group", "")).lower()
        ]
    if group.strip():
        contacts = [c for c in contacts if str(c.get("group", "")) == group.strip()]

    groups = sorted({str(c.get("group") or DEFAULT_GROUP) for c in load_contacts()})
    return {"ok": True, "contacts": contacts, "groups": groups, "total": len(contacts)}


@app.post("/api/contacts")
async def api_add_contacts(
    payload: ContactsIn, lang: str = Depends(request_language)
) -> dict[str, Any]:
    rows: list[dict[str, str]] = []
    if payload.contacts:
        rows = [c.model_dump() for c in payload.contacts]
    elif payload.email:
        rows = [{"name": payload.name, "email": payload.email, "group": payload.group}]

    if not rows:
        raise HTTPException(status_code=400, detail=tr(lang, "contact_email_required"))

    result = add_contacts(rows, lang)
    if not result["added"]:
        if result["duplicates"]:
            raise HTTPException(status_code=409, detail=tr(lang, "contact_exists"))
        raise HTTPException(status_code=400, detail=tr(lang, "contact_email_invalid"))
    return result


def parse_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        # Excel francophone exporte souvent en point-virgule.
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return list_from_rows([list(row) for row in reader])


def parse_xlsx(content: bytes, lang: str = DEFAULT_LANGUAGE) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(
            status_code=500, detail=tr(lang, "import_openpyxl_missing")
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail=tr(lang, "import_excel_unreadable", error=exc)
        ) from exc
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        raise HTTPException(status_code=400, detail=tr(lang, "import_no_sheet"))
    rows = [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return list_from_rows(rows)


def list_from_rows(rows: list[list[str]]) -> list[dict[str, str]]:
    """Transforme des lignes brutes en contacts, en détectant les en-têtes FR/EN."""
    rows = [r for r in rows if any(str(cell).strip() for cell in r)]
    if not rows:
        return []

    header_aliases = {
        "name": {"name", "nom", "prenom", "prénom", "fullname", "full name", "nom complet", "contact"},
        "email": {"email", "e-mail", "mail", "courriel", "adresse", "adresse email", "adresse e-mail"},
        "group": {"group", "groupe", "categorie", "catégorie", "category", "tag", "liste", "segment"},
    }
    first = [str(cell).strip().lower() for cell in rows[0]]
    mapping: dict[str, int] = {}
    for index, cell in enumerate(first):
        for field, aliases in header_aliases.items():
            if cell in aliases and field not in mapping:
                mapping[field] = index

    body = rows[1:] if mapping else rows
    if not mapping:
        # Pas d'en-tête reconnu : on devine la colonne contenant les adresses.
        email_col = 0
        for index in range(max(len(r) for r in rows)):
            if any(is_valid_email(str(r[index])) for r in rows if index < len(r)):
                email_col = index
                break
        mapping = {"email": email_col, "name": 1 if email_col == 0 else 0, "group": 2}

    out: list[dict[str, str]] = []
    for row in body:
        def cell(field: str) -> str:
            index = mapping.get(field, -1)
            return str(row[index]).strip() if 0 <= index < len(row) else ""

        email_addr = cell("email")
        if not email_addr:
            continue
        out.append({"name": cell("name"), "email": email_addr, "group": cell("group")})
    return out


@app.post("/api/contacts/import")
async def api_import_contacts(
    file: UploadFile = File(...), lang: str = Depends(request_language)
) -> dict[str, Any]:
    """Import de contacts depuis un fichier CSV ou XLSX (colonnes : name, email, group)."""
    if not file.filename:
        raise HTTPException(status_code=400, detail=tr(lang, "no_file"))
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail=tr(lang, "empty_file"))

    suffix = Path(file.filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx"}:
        rows = parse_xlsx(content, lang)
    elif suffix in {".csv", ".txt", ".tsv"}:
        rows = parse_csv(content)
    else:
        raise HTTPException(status_code=400, detail=tr(lang, "import_unsupported"))

    if not rows:
        raise HTTPException(status_code=400, detail=tr(lang, "import_no_contacts"))

    result = add_contacts(rows, lang)
    result["message"] = tr(lang, "import_done", summary=result["message"])
    result["ok"] = True
    return result


@app.delete("/api/contacts/{contact_id}")
async def api_delete_contact(
    contact_id: str, lang: str = Depends(request_language)
) -> dict[str, Any]:
    contacts = load_contacts()
    remaining = [c for c in contacts if str(c.get("id")) != contact_id]
    if len(remaining) == len(contacts):
        raise HTTPException(status_code=404, detail=tr(lang, "contact_not_found"))
    save_contacts(remaining)
    return {"ok": True, "message": tr(lang, "contact_deleted"), "contacts": remaining}


# --------------------------------- Modèles --------------------------------- #


@app.get("/api/templates")
async def api_templates(lang: str = Depends(request_language)) -> dict[str, Any]:
    templates = []
    for name in available_templates():
        meta = template_meta(lang, name)
        templates.append(
            {
                "name": name,
                "label": meta["label"],
                "description": meta["description"],
                "preview_url": f"/api/templates/{name}/preview",
            }
        )
    return {"ok": True, "templates": templates, "language": normalize_language(lang)}


@app.get("/api/templates/{name}/preview", response_class=HTMLResponse)
async def api_template_preview(
    name: str,
    accent: str = DEFAULT_ACCENT,
    lang: str = Depends(request_language),
) -> HTMLResponse:
    """Aperçu du modèle avec des données d'exemple, dans la langue de l'interface."""
    variables = sample_variables(lang)
    variables["accent_color"] = normalize_accent(accent)
    return HTMLResponse(render_email_html(name, variables, lang=lang))


@app.post("/api/preview", response_class=HTMLResponse)
async def api_preview(
    payload: PreviewRequest, lang: str = Depends(request_language)
) -> HTMLResponse:
    """Aperçu en direct : rend le modèle avec les variables saisies par l'utilisateur.

    Seuls les textes laissés vides sont remplacés par l'exemple traduit : ce que
    l'utilisateur écrit est affiché tel quel, dans la langue de son choix.
    """
    variables = payload.variables.model_dump()
    samples = sample_variables(lang)
    if not str(variables.get("title") or "").strip():
        variables["title"] = samples["title"]
    if not str(variables.get("body") or "").strip():
        variables["body"] = samples["body"]
    # Logo intégré en base64 : l'aperçu est autonome, exactement comme l'e-mail reçu.
    raw_logo = str(variables.get("logo_url") or "").strip()
    src = logo_preview_src(raw_logo)
    if raw_logo and not raw_logo.startswith(("http://", "https://", "data:")):
        try:
            src = logo_data_uri(resolve_upload(raw_logo, lang))
        except (HTTPException, OSError):
            src = ""
    return HTMLResponse(render_email_html(payload.template, variables, logo_src=src, lang=lang))


# ------------------------------- Paramètres -------------------------------- #


def public_config(config: dict[str, Any]) -> dict[str, Any]:
    """Configuration renvoyée au navigateur : mot de passe masqué."""
    return {
        "smtp_server": config["smtp_server"],
        "smtp_port": config["smtp_port"],
        "email": config["email"],
        "display_name": config["display_name"],
        "language": config["language"],
        "ui_accent_color": config["ui_accent_color"],
        "logo_url": config["logo_url"],
        "password": MASKED_PASSWORD if config["password"] else "",
        "has_password": bool(config["password"]),
    }


@app.get("/api/config")
async def api_get_config() -> dict[str, Any]:
    config = load_config()
    return {
        "ok": True,
        "config": public_config(config),
        "languages": [{"code": code, "label": LANGUAGE_LABELS[code]} for code in LANGUAGES],
        "configured": bool(config["smtp_server"] and config["email"] and config["password"]),
    }


def merge_config(payload: ConfigIn, lang: str = DEFAULT_LANGUAGE) -> dict[str, Any]:
    """Fusionne la saisie avec la configuration existante (mot de passe masqué conservé)."""
    current = load_config()
    try:
        port = int(str(payload.smtp_port).strip() or 587)
    except ValueError:
        raise HTTPException(status_code=400, detail=tr(lang, "port_not_number")) from None
    if not 1 <= port <= 65535:
        raise HTTPException(status_code=400, detail=tr(lang, "port_out_of_range"))

    password = payload.password or ""
    if not password or password == MASKED_PASSWORD:
        encoded = current["password"]
    else:
        encoded = encode_password(password)

    # Ni la langue ni la couleur de l'interface ne sont modifiées par le formulaire
    # SMTP : on conserve les valeurs enregistrées, sauf indication explicite.
    language = normalize_language(payload.language or current["language"])
    ui_accent = normalize_accent(payload.ui_accent_color or current["ui_accent_color"])

    return {
        "smtp_server": payload.smtp_server.strip(),
        "smtp_port": port,
        "email": payload.email.strip(),
        "password": encoded,
        "display_name": payload.display_name.strip(),
        "language": language,
        "ui_accent_color": ui_accent,
        "logo_url": current["logo_url"],
    }


@app.post("/api/config")
async def api_save_config(
    payload: ConfigIn, lang: str = Depends(request_language)
) -> dict[str, Any]:
    config = merge_config(payload, lang)
    if not config["smtp_server"]:
        raise HTTPException(status_code=400, detail=tr(lang, "server_required"))
    if not is_valid_email(config["email"]):
        raise HTTPException(status_code=400, detail=tr(lang, "config_email_invalid"))
    write_json(CONFIG_FILE, config)
    return {
        "ok": True,
        "message": tr(lang, "config_saved"),
        "config": public_config(config),
    }


@app.post("/api/language")
async def api_save_language(payload: LanguageIn) -> dict[str, Any]:
    """Enregistre la langue de l'interface, sans toucher au reste de la configuration.

    Endpoint distinct de /api/config : changer de langue ne doit pas exiger une
    configuration SMTP complète.
    """
    requested = str(payload.language or "").strip()
    language = normalize_language(requested)
    if requested and language != requested.lower().replace("_", "-").split("-")[0]:
        raise HTTPException(
            status_code=400,
            detail=tr(
                language, "language_unknown", language=requested, allowed=", ".join(LANGUAGES)
            ),
        )

    config = load_config()
    config["language"] = language
    write_json(CONFIG_FILE, config)
    return {
        "ok": True,
        "message": tr(language, "language_saved"),
        "language": language,
        "config": public_config(config),
    }


@app.post("/api/theme")
async def api_save_theme(
    payload: ThemeIn, lang: str = Depends(request_language)
) -> dict[str, Any]:
    """Enregistre la couleur d'accentuation de l'interface.

    Endpoint distinct de /api/config, comme /api/language : changer de thème ne
    doit pas exiger une configuration SMTP complète.
    """
    requested = str(payload.ui_accent_color or "").strip()
    color = normalize_accent(requested)
    if requested and color.lower() != requested.lower():
        raise HTTPException(status_code=400, detail=tr(lang, "theme_invalid", color=requested))

    config = load_config()
    config["ui_accent_color"] = color
    write_json(CONFIG_FILE, config)
    return {
        "ok": True,
        "message": tr(lang, "theme_saved"),
        "ui_accent_color": color,
        "config": public_config(config),
    }


@app.post("/api/logo")
async def api_save_logo(
    payload: LogoIn, lang: str = Depends(request_language)
) -> dict[str, Any]:
    """Mémorise la dernière adresse de logo utilisée en mode URL.

    Comme /api/language et /api/theme : un réglage de confort, sans rapport avec
    la configuration SMTP, donc son propre endpoint.
    """
    url = normalize_logo_url(payload.logo_url)
    if url is None:
        raise HTTPException(
            status_code=400, detail=tr(lang, "logo_url_invalid", url=payload.logo_url)
        )

    config = load_config()
    config["logo_url"] = url
    write_json(CONFIG_FILE, config)
    return {
        "ok": True,
        "message": tr(lang, "logo_saved"),
        "logo_url": url,
        "config": public_config(config),
    }


@app.post("/api/config/test")
async def api_test_config(
    payload: ConfigIn | None = Body(default=None),
    lang: str = Depends(request_language),
) -> dict[str, Any]:
    """Teste la connexion SMTP avec la saisie en cours (ou la configuration enregistrée)."""
    config = merge_config(payload, lang) if payload is not None else load_config()
    return await asyncio.to_thread(lambda: test_smtp_connection(config, lang))


# --------------------------------------------------------------------------- #
# Lancement
# --------------------------------------------------------------------------- #

def startup_banner() -> None:
    """Affiche l'URL de l'application (repli sans emoji si la console ne gère pas l'UTF-8)."""
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # consoles Windows
    except Exception:
        pass
    message = f"🚀 Zm@il running at http://localhost:{PORT}"
    try:
        print(message, flush=True)
    except UnicodeEncodeError:
        print(message.encode("ascii", "ignore").decode("ascii").strip(), flush=True)


if __name__ == "__main__":
    import uvicorn

    ensure_directories()
    startup_banner()
    uvicorn.run(app, host=HOST, port=PORT, log_level="info")
